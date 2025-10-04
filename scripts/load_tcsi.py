#!/usr/bin/env python3
"""Builds and loads the local TCSI SQLite database from the supplied extracts."""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import re
import sqlite3
import tempfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple

try:  # Optional dependency for Excel ingestion
    import openpyxl  # type: ignore
except ModuleNotFoundError:  # pragma: no cover - handled at runtime
    openpyxl = None

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DB_PATH = ROOT / "tcsi.db"
SCHEMA_PATH = ROOT / "database" / "schema.sql"

SUPPORTED_TABULAR_SUFFIXES = {".csv", ".xlsx", ".xlsm"}

STAGING_NAME_OVERRIDES = {
    "HEPStudents": "stg_hep_students",
    "HEPStudentCitizenships": "stg_hep_student_citizenships",
    "HEPStudentDisabilities": "stg_hep_student_disabilities",
    "StudentContactsFirstReportedAddress": "stg_student_contacts_first_reported_address",
    "CoursesOfStudy": "stg_courses_of_study",
    "HEPCourses": "stg_hep_courses",
    "HEPCoursesOnCampuses": "stg_hep_courses_on_campuses",
    "CampusesTAC": "stg_campuses_tac",
    "CampusCourseFeesITSP": "stg_campus_course_fees_itsp",
    "CourseFieldsOfEducation": "stg_course_fields_of_education",
    "CourseSpecialisations": "stg_course_specialisations",
    "HEPCourseAdmissions": "stg_hep_course_admissions",
    "HEPCoursePriorCredits": "stg_hep_course_prior_credits",
    "HEP_units-AOUs": "stg_hep_units_aous",
    "OSHELP": "stg_oshelp",
    "SAHELP": "stg_sahelp",
    "CommonwealthScholarships": "stg_commonwealth_scholarships",
    "RTPScholarships": "stg_rtp_scholarships",
    "RTPStipend": "stg_rtp_stipend",
    "AggregatedAwards": "stg_aggregated_awards",
    "HEPHDREndUsersEngagement": "stg_hephdr_end_users_engagement",
    "UndeterminedStudentsRecords": "stg_undetermined_students_records",
    "UndeterminedStudentCitizenships": "stg_undetermined_student_citizenships",
    "UndeterminedStudentDisabilities": "stg_undetermined_student_disabilities",
    "SpecialInterestCourses": "stg_special_interest_courses",
}

_COLUMN_PATTERN = re.compile(r"\{\{([^}]+)\}\}")


def sanitize_identifier(name: str) -> str:
    cleaned = name.replace("-", "_").replace(" ", "_")
    cleaned = re.sub(r"[^0-9A-Za-z_]+", "_", cleaned)
    cleaned = re.sub(r"__+", "_", cleaned)
    cleaned = cleaned.strip("_")
    cleaned = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", cleaned)
    return cleaned.lower()


def render_sql(template: str) -> str:
    return _COLUMN_PATTERN.sub(lambda m: sanitize_identifier(m.group(1)), template)


def to_snake(name: str) -> str:
    return sanitize_identifier(name)


def sanitize_column_name(name: str, existing: Optional[set] = None) -> str:
    base = sanitize_identifier(name)
    existing = existing or set()
    candidate = base or "col"
    counter = 1
    while candidate in existing:
        counter += 1
        candidate = f"{base}_{counter}"
    existing.add(candidate)
    return candidate


def parse_entity_and_year(stem: str) -> Tuple[str, Optional[str]]:
    cleaned = re.sub(r"^deidentified_", "", stem, flags=re.IGNORECASE)
    cleaned = re.sub(r"^\d+_tcsi_DataExtract_\d+_11_", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^\d+_", "", cleaned)
    match = re.match(r"(.+?)_((?:19|20)\d{2})(-onwards)?$", cleaned)
    if match:
        base, year, onwards = match.groups()
        year_token = f"{year}{'-onwards' if onwards else ''}"
        return base, year_token
    return cleaned, None


def staging_table_for(entity: str) -> str:
    return STAGING_NAME_OVERRIDES.get(entity, f"stg_{to_snake(entity)}")


def ensure_staging_table(cur: sqlite3.Cursor, table: str, columns: List[str]) -> None:
    cols_sql = ", \n        ".join(f"{col} TEXT" for col in columns)
    cur.execute(
        f"""
        CREATE TABLE IF NOT EXISTS {table} (
            {cols_sql},
            source_file TEXT NOT NULL,
            reporting_year TEXT,
            extraction_timestamp TEXT NOT NULL
        )
        """
    )


def read_csv_header(csv_path: Path) -> List[str]:
    with csv_path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        try:
            header = next(reader)
        except StopIteration as exc:
            raise ValueError(f"CSV {csv_path} has no header") from exc
    return header


def load_csv(
    conn: sqlite3.Connection,
    csv_path: Path,
    extraction_ts: str,
    *,
    source_name: Optional[str] = None,
    stem_override: Optional[str] = None,
) -> Tuple[str, int]:
    stem = stem_override or csv_path.stem
    entity, year_token = parse_entity_and_year(stem)
    table = staging_table_for(entity)

    raw_header = read_csv_header(csv_path)
    existing: set = set()
    sanitized_columns = [sanitize_column_name(col, existing) for col in raw_header]

    cur = conn.cursor()
    ensure_staging_table(cur, table, sanitized_columns)
    cur.execute(
        f"DELETE FROM {table} WHERE extraction_timestamp = ? AND source_file = ?",
        (extraction_ts, source_name or csv_path.name),
    )

    placeholders = ", ".join(["?"] * (len(sanitized_columns) + 3))
    insert_sql = f"INSERT INTO {table} ({', '.join(sanitized_columns)}, source_file, reporting_year, extraction_timestamp) VALUES ({placeholders})"

    row_count = 0
    with csv_path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        batch: List[Tuple[str, ...]] = []
        for raw_row in reader:
            values = [raw_row.get(col, "").strip() for col in raw_header]
            batch.append(tuple(values + [source_name or csv_path.name, year_token, extraction_ts]))
            row_count += 1
            if len(batch) >= 1000:
                cur.executemany(insert_sql, batch)
                batch.clear()
        if batch:
            cur.executemany(insert_sql, batch)
    conn.commit()
    return table, row_count


def load_excel(conn: sqlite3.Connection, excel_path: Path, extraction_ts: str) -> Tuple[str, int]:
    if openpyxl is None:
        raise RuntimeError(
            "Excel ingestion requires the optional dependency 'openpyxl'. "
            "Install it via `pip install openpyxl`."
        )

    workbook = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as exc:  # pragma: no cover - defensive guard
            raise ValueError(f"Excel file {excel_path} has no header row") from exc
        raw_header = ["" if cell is None else str(cell) for cell in header_row]
        # Write to a temporary CSV using the same stem logic for downstream processing.
        with tempfile.NamedTemporaryFile(
            mode="w",
            newline="",
            suffix=".csv",
            prefix=f"{excel_path.stem}_",
            delete=False,
            encoding="utf-8",
        ) as tmp:
            writer = csv.writer(tmp)
            writer.writerow(raw_header)
            header_len = len(raw_header)
            for row in rows:
                values = []
                for idx in range(header_len):
                    cell = row[idx] if idx < len(row) else None
                    values.append("" if cell is None else str(cell))
                writer.writerow(values)
            temp_path = Path(tmp.name)
    finally:
        workbook.close()

    try:
        return load_csv(
            conn,
            temp_path,
            extraction_ts,
            source_name=excel_path.name,
            stem_override=excel_path.stem,
        )
    finally:
        temp_path.unlink(missing_ok=True)


def apply_schema(conn: sqlite3.Connection) -> None:
    schema_sql = SCHEMA_PATH.read_text(encoding="utf-8")
    conn.executescript(schema_sql)
    conn.commit()


def staging_table_exists(conn: sqlite3.Connection, table: str) -> bool:
    row = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = ?",
        (table,),
    ).fetchone()
    return row is not None


def staging_has_column(conn: sqlite3.Connection, table: str, original_name: str) -> bool:
    sanitized = sanitize_identifier(original_name)
    cur = conn.execute(f"PRAGMA table_info({table})")
    return any(row[1] == sanitized for row in cur.fetchall())


def placeholder_or_null(conn: sqlite3.Connection, table: str, original_name: str, default: str = "NULL") -> str:
    return f"{{{{{original_name}}}}}" if staging_has_column(conn, table, original_name) else default


def nullif_or_null(conn: sqlite3.Connection, table: str, original_name: str) -> str:
    return f"NULLIF({{{{{original_name}}}}}, 'NULL')" if staging_has_column(conn, table, original_name) else "NULL"


def upsert_dim_provider(conn: sqlite3.Connection, extraction_ts: str) -> None:
    if not staging_table_exists(conn, "stg_hep_students"):
        return
    sql = render_sql(
        """
        INSERT INTO dim_provider (provider_res_key, provider_code, provider_type)
        SELECT DISTINCT {{UID1_ProvidersResKey}}, {{E306_ProviderCode}}, {{E781_ProviderType}}
        FROM stg_hep_students
        WHERE extraction_timestamp = ?
          AND {{UID1_ProvidersResKey}} IS NOT NULL
        ON CONFLICT(provider_res_key) DO UPDATE SET
            provider_code = excluded.provider_code,
            provider_type = excluded.provider_type
        """
    )
    conn.execute(sql, (extraction_ts,))


def upsert_dim_student(conn: sqlite3.Connection, extraction_ts: str) -> None:
    if not staging_table_exists(conn, "stg_hep_students"):
        return
    sql = render_sql(
        """
        INSERT INTO dim_student (
            student_res_key, provider_res_key, student_identifier, chessn, usi,
            usi_verification_status, tfn_verification_status, date_of_birth,
            family_name, given_name_first, given_name_others,
            residential_address_line1, residential_address_suburb,
            residential_address_postcode, residential_address_state,
            residential_address_country, term_residence_postcode,
            term_residence_country, gender_code, atsi_code,
            country_of_birth_code, arrival_year, language_home_code,
            year_left_school, level_left_school,
            highest_ed_parent1, highest_ed_parent2, extraction_timestamp
        )
        SELECT
            {{UID8_StudentsResKey}},
            {{UID1_ProvidersResKey}},
            {{E313_StudentIdentificationCode}},
            {{E488_CHESSN}},
            {{E584_USI}},
            {{A170_USIVerificationStatus}},
            {{A167_TFNVerificationStatus}},
            {{E314_DateOfBirth}},
            {{E402_StudentFamilyName}},
            {{E403_StudentGivenNameFirst}},
            {{E404_StudentGivenNameOthers}},
            {{E410_ResidentialAddressLine1}},
            {{E469_ResidentialAddressSuburb}},
            {{E320_ResidentialAddressPostcode}},
            {{E470_ResidentialAddressState}},
            {{E658_ResidentialAddressCountryCode}},
            {{E319_TermResidencePostcode}},
            {{E661_TermResidenceCountryCode}},
            {{E315_GenderCode}},
            {{E316_ATSICode}},
            {{E346_CountryOfBirthCode}},
            {{E347_ArrivalInAustraliaYear}},
            {{E348_LanguageSpokenAtHomeCode}},
            {{E572_YearLeftSchool}},
            {{E612_LevelLeftSchool}},
            {{E573_HighestEducationParent1}},
            {{E574_HighestEducationParent2}},
            ?
        FROM stg_hep_students
        WHERE extraction_timestamp = ?
          AND {{UID8_StudentsResKey}} IS NOT NULL
        ON CONFLICT(student_res_key) DO UPDATE SET
            provider_res_key = excluded.provider_res_key,
            student_identifier = excluded.student_identifier,
            chessn = excluded.chessn,
            usi = excluded.usi,
            usi_verification_status = excluded.usi_verification_status,
            tfn_verification_status = excluded.tfn_verification_status,
            date_of_birth = excluded.date_of_birth,
            family_name = excluded.family_name,
            given_name_first = excluded.given_name_first,
            given_name_others = excluded.given_name_others,
            residential_address_line1 = excluded.residential_address_line1,
            residential_address_suburb = excluded.residential_address_suburb,
            residential_address_postcode = excluded.residential_address_postcode,
            residential_address_state = excluded.residential_address_state,
            residential_address_country = excluded.residential_address_country,
            term_residence_postcode = excluded.term_residence_postcode,
            term_residence_country = excluded.term_residence_country,
            gender_code = excluded.gender_code,
            atsi_code = excluded.atsi_code,
            country_of_birth_code = excluded.country_of_birth_code,
            arrival_year = excluded.arrival_year,
            language_home_code = excluded.language_home_code,
            year_left_school = excluded.year_left_school,
            level_left_school = excluded.level_left_school,
            highest_ed_parent1 = excluded.highest_ed_parent1,
            highest_ed_parent2 = excluded.highest_ed_parent2,
            extraction_timestamp = excluded.extraction_timestamp
        """
    )
    conn.execute(sql, (extraction_ts, extraction_ts))


def upsert_student_contact(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_student_contacts_first_reported_address"
    if not staging_table_exists(conn, table):
        return
    sql = render_sql(
        """
        SELECT
            {{UID8_StudentsResKey}},
            {{UID1_ProvidersResKey}},
            {{E313_StudentIdentificationCode}},
            {{E787_FirstResidentialAddressLine1}},
            {{E789_FirstResidentialAddressSuburb}},
            {{E791_FirstResidentialAddressState}},
            {{E659_FirstResidentialAddressCountryCode}},
            {{E790_FirstResidentialAddressPostcode}}
        FROM stg_student_contacts_first_reported_address
        WHERE extraction_timestamp = ?
        """
    )
    cur = conn.cursor()
    rows = cur.execute(sql, (extraction_ts,)).fetchall()
    insert_sql = """
        INSERT INTO bridge_student_contact (
            student_contact_res_key, student_res_key, provider_res_key,
            student_identifier, first_address_line1, first_address_suburb,
            first_address_state, first_address_country, first_address_postcode,
            extraction_timestamp
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(student_contact_res_key) DO UPDATE SET
            provider_res_key = excluded.provider_res_key,
            student_identifier = excluded.student_identifier,
            first_address_line1 = excluded.first_address_line1,
            first_address_suburb = excluded.first_address_suburb,
            first_address_state = excluded.first_address_state,
            first_address_country = excluded.first_address_country,
            first_address_postcode = excluded.first_address_postcode,
            extraction_timestamp = excluded.extraction_timestamp
    """
    payload = []
    for row in rows:
        student_res_key = row[0]
        provider_res_key = row[1]
        if not student_res_key:
            continue
        hash_input = "|".join(str(part) for part in row + (extraction_ts,))
        res_key = hashlib.sha1(hash_input.encode("utf-8")).hexdigest()
        payload.append((
            res_key,
            student_res_key,
            provider_res_key,
            row[2],
            row[3],
            row[4],
            row[5],
            row[6],
            row[7],
            extraction_ts,
        ))
    if payload:
        cur.executemany(insert_sql, payload)
        conn.commit()


def upsert_student_citizenship(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_hep_student_citizenships"
    if not staging_table_exists(conn, table):
        return
    sql = render_sql(
        """
        INSERT INTO bridge_student_citizenship (
            student_citizenship_res_key, student_res_key, provider_res_key,
            student_identifier, chessn, citizen_resident_code,
            effective_from_date, effective_to_date, extraction_timestamp
        )
        SELECT
            {{UID10_StudentCitizenshipsResKey}},
            {{UID8_StudentsResKey}},
            {{UID1_ProvidersResKey}},
            {{E313_StudentIdentificationCode}},
            {{E488_CHESSN}},
            {{E358_CitizenResidentCode}},
            {{E609_EffectiveFromDate}},
            {{E610_EffectiveToDate}},
            ?
        FROM stg_hep_student_citizenships
        WHERE extraction_timestamp = ?
          AND {{UID10_StudentCitizenshipsResKey}} IS NOT NULL
        ON CONFLICT(student_citizenship_res_key) DO UPDATE SET
            student_res_key = excluded.student_res_key,
            provider_res_key = excluded.provider_res_key,
            student_identifier = excluded.student_identifier,
            chessn = excluded.chessn,
            citizen_resident_code = excluded.citizen_resident_code,
            effective_from_date = excluded.effective_from_date,
            effective_to_date = excluded.effective_to_date,
            extraction_timestamp = excluded.extraction_timestamp
        """
    )
    conn.execute(sql, (extraction_ts, extraction_ts))


def upsert_student_disability(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_hep_student_disabilities"
    if not staging_table_exists(conn, table):
        return
    sql = render_sql(
        """
        INSERT INTO bridge_student_disability (
            student_disability_res_key, student_res_key, provider_res_key,
            student_identifier, chessn, disability_code,
            effective_from_date, effective_to_date, extraction_timestamp
        )
        SELECT
            {{UID11_StudentDisabilitiesResKey}},
            {{UID8_StudentsResKey}},
            {{UID1_ProvidersResKey}},
            {{E313_StudentIdentificationCode}},
            {{E488_CHESSN}},
            {{E615_DisabilityCode}},
            {{E609_EffectiveFromDate}},
            {{E610_EffectiveToDate}},
            ?
        FROM stg_hep_student_disabilities
        WHERE extraction_timestamp = ?
          AND {{UID11_StudentDisabilitiesResKey}} IS NOT NULL
        ON CONFLICT(student_disability_res_key) DO UPDATE SET
            student_res_key = excluded.student_res_key,
            provider_res_key = excluded.provider_res_key,
            student_identifier = excluded.student_identifier,
            chessn = excluded.chessn,
            disability_code = excluded.disability_code,
            effective_from_date = excluded.effective_from_date,
            effective_to_date = excluded.effective_to_date,
            extraction_timestamp = excluded.extraction_timestamp
        """
    )
    conn.execute(sql, (extraction_ts, extraction_ts))


def upsert_course_of_study(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_courses_of_study"
    if not staging_table_exists(conn, table):
        return
    sql = render_sql(
        """
        INSERT INTO dim_course_of_study (
            course_of_study_res_key, provider_res_key, course_of_study_code,
            course_of_study_name, course_of_study_type, course_of_study_load,
            combined_course_indicator
        )
        SELECT DISTINCT
            {{UID3_CoursesOfStudyResKey}},
            {{UID1_ProvidersResKey}},
            {{E533_CourseOfStudyCode}},
            {{E394_CourseOfStudyName}},
            {{E310_CourseOfStudyType}},
            {{E350_CourseOfStudyLoad}},
            {{E455_CombinedCourseOfStudyIndicator}}
        FROM stg_courses_of_study
        WHERE extraction_timestamp = ?
          AND {{UID3_CoursesOfStudyResKey}} IS NOT NULL
        ON CONFLICT(course_of_study_res_key) DO UPDATE SET
            provider_res_key = excluded.provider_res_key,
            course_of_study_code = excluded.course_of_study_code,
            course_of_study_name = excluded.course_of_study_name,
            course_of_study_type = excluded.course_of_study_type,
            course_of_study_load = excluded.course_of_study_load,
            combined_course_indicator = excluded.combined_course_indicator
        """
    )
    conn.execute(sql, (extraction_ts,))


def upsert_course(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_hep_courses"
    if not staging_table_exists(conn, table):
        return
    template = """
        INSERT INTO dim_course (
            course_res_key, provider_res_key, course_of_study_res_key,
            course_code, course_name, standard_course_duration,
            cricos_code, effective_from_date, effective_to_date
        )
        SELECT DISTINCT
            {{UID5_CoursesResKey}},
            {{UID1_ProvidersResKey}},
            {{UID3_CoursesOfStudyResKey}},
            {{E307_CourseCode}},
            {{E308_CourseName}},
            {{E596_StandardCourseDuration}},
            __CRICOS__,
            {{E609_EffectiveFromDate}},
            {{E610_EffectiveToDate}}
        FROM stg_hep_courses
        WHERE extraction_timestamp = ?
          AND {{UID5_CoursesResKey}} IS NOT NULL
        ON CONFLICT(course_res_key) DO UPDATE SET
            provider_res_key = excluded.provider_res_key,
            course_of_study_res_key = excluded.course_of_study_res_key,
            course_code = excluded.course_code,
            course_name = excluded.course_name,
            standard_course_duration = excluded.standard_course_duration,
            cricos_code = excluded.cricos_code,
            effective_from_date = excluded.effective_from_date,
            effective_to_date = excluded.effective_to_date
        """
    cricos_expr = "{{E597_CRICOSCode}}" if staging_has_column(conn, table, "E597_CRICOSCode") else "NULL"
    sql = render_sql(template.replace("__CRICOS__", cricos_expr))
    conn.execute(sql, (extraction_ts,))


def upsert_course_campus(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_hep_courses_on_campuses"
    if not staging_table_exists(conn, table):
        return
    campus_sql = render_sql(
        """
        INSERT INTO dim_campus (
            campus_res_key, provider_res_key, campus_suburb,
            campus_country_code, campus_postcode,
            first_effective_from_date, last_effective_to_date
        )
        SELECT
            {{UID2_CampusesResKey}},
            {{UID1_ProvidersResKey}},
            {{E525_CampusSuburb}},
            {{E644_CampusCountryCode}},
            {{E559_CampusPostcode}},
            MIN(NULLIF({{Campuses_E609_EffectiveFromDate}}, 'NULL')),
            MAX(NULLIF({{Campuses_E610_EffectiveToDate}}, 'NULL'))
        FROM stg_hep_courses_on_campuses
        WHERE extraction_timestamp = ?
          AND {{UID2_CampusesResKey}} IS NOT NULL
        GROUP BY {{UID2_CampusesResKey}}, {{UID1_ProvidersResKey}}
        ON CONFLICT(campus_res_key) DO UPDATE SET
            provider_res_key = excluded.provider_res_key,
            campus_suburb = excluded.campus_suburb,
            campus_country_code = excluded.campus_country_code,
            campus_postcode = excluded.campus_postcode,
            first_effective_from_date = excluded.first_effective_from_date,
            last_effective_to_date = excluded.last_effective_to_date
        """
    )
    conn.execute(campus_sql, (extraction_ts,))
    sql = render_sql(
        """
        INSERT INTO bridge_course_campus (
            course_campus_res_key, provider_res_key, course_res_key, campus_res_key,
            campus_suburb, campus_country_code, campus_postcode,
            campus_effective_from, campus_effective_to,
            course_effective_from, course_effective_to, extraction_timestamp
        )
        SELECT
            {{UID4_CoursesOnCampusResKey}},
            {{UID1_ProvidersResKey}},
            {{UID5_CoursesResKey}},
            {{UID2_CampusesResKey}},
            {{E525_CampusSuburb}},
            {{E644_CampusCountryCode}},
            {{E559_CampusPostcode}},
            {{Campuses_E609_EffectiveFromDate}},
            {{Campuses_E610_EffectiveToDate}},
            {{E609_EffectiveFromDate}},
            {{E610_EffectiveToDate}},
            ?
        FROM stg_hep_courses_on_campuses
        WHERE extraction_timestamp = ?
          AND {{UID4_CoursesOnCampusResKey}} IS NOT NULL
          AND {{UID5_CoursesResKey}} IS NOT NULL
          AND EXISTS (
              SELECT 1 FROM dim_course dc WHERE dc.course_res_key = {{UID5_CoursesResKey}}
          )
        ON CONFLICT(course_campus_res_key) DO UPDATE SET
            provider_res_key = excluded.provider_res_key,
            course_res_key = excluded.course_res_key,
            campus_res_key = excluded.campus_res_key,
            campus_suburb = excluded.campus_suburb,
            campus_country_code = excluded.campus_country_code,
            campus_postcode = excluded.campus_postcode,
            campus_effective_from = excluded.campus_effective_from,
            campus_effective_to = excluded.campus_effective_to,
            course_effective_from = excluded.course_effective_from,
            course_effective_to = excluded.course_effective_to,
            extraction_timestamp = excluded.extraction_timestamp
        """
    )
    conn.execute(sql, (extraction_ts, extraction_ts))

    conn.execute(
        """
        INSERT INTO dim_campus (
            campus_res_key, provider_res_key, campus_suburb,
            campus_country_code, campus_postcode,
            first_effective_from_date, last_effective_to_date
        )
        SELECT
            campus_res_key,
            provider_res_key,
            MIN(campus_suburb),
            MIN(campus_country_code),
            MIN(campus_postcode),
            MIN(NULLIF(campus_effective_from, 'NULL')),
            MAX(NULLIF(campus_effective_to, 'NULL'))
        FROM bridge_course_campus
        GROUP BY campus_res_key, provider_res_key
        ON CONFLICT(campus_res_key) DO UPDATE SET
            provider_res_key = excluded.provider_res_key,
            campus_suburb = excluded.campus_suburb,
            campus_country_code = excluded.campus_country_code,
            campus_postcode = excluded.campus_postcode,
            first_effective_from_date = excluded.first_effective_from_date,
            last_effective_to_date = excluded.last_effective_to_date
        """
    )


def upsert_course_field_of_education(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_course_fields_of_education"
    if not staging_table_exists(conn, table):
        return
    sql = render_sql(
        """
        INSERT INTO bridge_course_field_of_education (
            course_field_res_key, provider_res_key, course_of_study_res_key,
            course_res_key, course_code, course_of_study_name,
            field_of_education_code, field_of_education_supp_code,
            effective_from_date, effective_to_date, extraction_timestamp
        )
        SELECT
            {{UID48_CourseFieldsOfEducationResKey}},
            {{UID1_ProvidersResKey}},
            {{UID3_CoursesOfStudyResKey}},
            {{UID5_CoursesResKey}},
            {{E307_CourseCode}},
            {{E394_CourseOfStudyName}},
            {{E461_FieldOfEducationCode}},
            {{E462_FieldOfEducationSupplementaryCode}},
            {{E609_EffectiveFromDate}},
            {{E610_EffectiveToDate}},
            ?
        FROM stg_course_fields_of_education
        WHERE extraction_timestamp = ?
          AND {{UID48_CourseFieldsOfEducationResKey}} IS NOT NULL
        ON CONFLICT(course_field_res_key) DO UPDATE SET
            provider_res_key = excluded.provider_res_key,
            course_of_study_res_key = excluded.course_of_study_res_key,
            course_res_key = excluded.course_res_key,
            course_code = excluded.course_code,
            course_of_study_name = excluded.course_of_study_name,
            field_of_education_code = excluded.field_of_education_code,
            field_of_education_supp_code = excluded.field_of_education_supp_code,
            effective_from_date = excluded.effective_from_date,
            effective_to_date = excluded.effective_to_date,
            extraction_timestamp = excluded.extraction_timestamp
        """
    )
    conn.execute(sql, (extraction_ts, extraction_ts))


def upsert_course_admission(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_hep_course_admissions"
    if not staging_table_exists(conn, table):
        return
    mode_expr = placeholder_or_null(conn, table, "E329_ModeOfAttendanceCode")
    atar_expr = nullif_or_null(conn, table, "E632_ATAR")
    selection_expr = nullif_or_null(conn, table, "E605_SelectionRank")
    template = """
        INSERT INTO fact_course_admission (
            course_admission_res_key, provider_res_key, student_res_key,
            course_res_key, student_identifier, course_code, course_name,
            course_of_study_code, course_of_study_type, course_commencement_date,
            attendance_type_code, mode_of_attendance_code,
            course_outcome_code, course_outcome_date, chessn,
            hdr_thesis_submission_date, atar, selection_rank,
            highest_attainment_code, hdr_primary_for_code, hdr_secondary_for_code,
            extraction_timestamp
        )
        SELECT
            {{UID15_CourseAdmissionsResKey}},
            {{UID1_ProvidersResKey}},
            {{UID8_StudentsResKey}},
            {{UID5_CoursesResKey}},
            {{E313_StudentIdentificationCode}},
            {{E307_CourseCode}},
            {{E308_CourseName}},
            {{E533_CourseOfStudyCode}},
            {{E310_CourseOfStudyType}},
            {{E534_CourseOfStudyCommencementDate}},
            {{E330_AttendanceTypeCode}},
            __MODE__,
            {{E599_CourseOutcomeCode}},
            {{E592_CourseOutcomeDate}},
            {{E488_CHESSN}},
            {{E591_HDRThesisSubmissionDate}},
            __ATAR__,
            __SELECTION__,
            {{E620_HighestAttainmentCode}},
            {{E594_HDRPrimaryFieldOfResearchCode}},
            {{E595_HDRSecondaryFieldOfResearchCode}},
            ?
        FROM stg_hep_course_admissions
        WHERE extraction_timestamp = ?
          AND {{UID15_CourseAdmissionsResKey}} IS NOT NULL
        ON CONFLICT(course_admission_res_key) DO UPDATE SET
            provider_res_key = excluded.provider_res_key,
            student_res_key = excluded.student_res_key,
            course_res_key = excluded.course_res_key,
            student_identifier = excluded.student_identifier,
            course_code = excluded.course_code,
            course_name = excluded.course_name,
            course_of_study_code = excluded.course_of_study_code,
            course_of_study_type = excluded.course_of_study_type,
            course_commencement_date = excluded.course_commencement_date,
            attendance_type_code = excluded.attendance_type_code,
            mode_of_attendance_code = excluded.mode_of_attendance_code,
            course_outcome_code = excluded.course_outcome_code,
            course_outcome_date = excluded.course_outcome_date,
            chessn = excluded.chessn,
            hdr_thesis_submission_date = excluded.hdr_thesis_submission_date,
            atar = excluded.atar,
            selection_rank = excluded.selection_rank,
            highest_attainment_code = excluded.highest_attainment_code,
            hdr_primary_for_code = excluded.hdr_primary_for_code,
            hdr_secondary_for_code = excluded.hdr_secondary_for_code,
            extraction_timestamp = excluded.extraction_timestamp
        """
    template = template.replace("__MODE__", mode_expr).replace("__ATAR__", atar_expr).replace("__SELECTION__", selection_expr)
    sql = render_sql(template)
    conn.execute(sql, (extraction_ts, extraction_ts))


def upsert_course_prior_credit(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_hep_course_prior_credits"
    if not staging_table_exists(conn, table):
        return
    sql = render_sql(
        """
        INSERT INTO fact_course_prior_credit (
            course_prior_credit_res_key, provider_res_key, course_admission_res_key,
            student_identifier, course_code, course_commencement_date,
            credit_used_value, credit_basis, credit_provider_code, extraction_timestamp
        )
        SELECT
            {{UID32_CoursePriorCreditsResKey}},
            {{UID1_ProvidersResKey}},
            {{UID15_CourseAdmissionsResKey}},
            {{E313_StudentIdentificationCode}},
            {{E307_CourseCode}},
            {{E534_CourseOfStudyCommencementDate}},
            NULLIF({{E560_CreditUsedValue}}, 'NULL'),
            {{E561_CreditBasis}},
            {{E566_CreditProviderCode}},
            ?
        FROM stg_hep_course_prior_credits
        WHERE extraction_timestamp = ?
          AND {{UID32_CoursePriorCreditsResKey}} IS NOT NULL
        ON CONFLICT(course_prior_credit_res_key) DO UPDATE SET
            provider_res_key = excluded.provider_res_key,
            course_admission_res_key = excluded.course_admission_res_key,
            student_identifier = excluded.student_identifier,
            course_code = excluded.course_code,
            course_commencement_date = excluded.course_commencement_date,
            credit_used_value = excluded.credit_used_value,
            credit_basis = excluded.credit_basis,
            credit_provider_code = excluded.credit_provider_code,
            extraction_timestamp = excluded.extraction_timestamp
        """
    )
    conn.execute(sql, (extraction_ts, extraction_ts))


def upsert_course_specialisation(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_course_specialisations"
    if not staging_table_exists(conn, table):
        return
    student_expr = placeholder_or_null(conn, table, "UID8_StudentsResKey")
    template = """
        INSERT INTO fact_course_specialisation (
            course_specialisation_res_key, provider_res_key, course_admission_res_key,
            student_res_key, student_identifier, course_code,
            course_commencement_date, specialisation_code, extraction_timestamp
        )
        SELECT
            {{UID33_CourseSpecialisationsResKey}},
            {{UID1_ProvidersResKey}},
            {{UID15_CourseAdmissionsResKey}},
            __STUDENT__,
            {{E313_StudentIdentificationCode}},
            {{E307_CourseCode}},
            {{E534_CourseOfStudyCommencementDate}},
            {{E463_SpecialisationCode}},
            ?
        FROM stg_course_specialisations
        WHERE extraction_timestamp = ?
          AND {{UID33_CourseSpecialisationsResKey}} IS NOT NULL
        ON CONFLICT(course_specialisation_res_key) DO UPDATE SET
            provider_res_key = excluded.provider_res_key,
            course_admission_res_key = excluded.course_admission_res_key,
            student_res_key = excluded.student_res_key,
            student_identifier = excluded.student_identifier,
            course_code = excluded.course_code,
            course_commencement_date = excluded.course_commencement_date,
            specialisation_code = excluded.specialisation_code,
            extraction_timestamp = excluded.extraction_timestamp
        """
    template = template.replace("__STUDENT__", student_expr)
    sql = render_sql(template)
    conn.execute(sql, (extraction_ts, extraction_ts))


def upsert_unit_enrolment(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_hep_units_aous"
    if not staging_table_exists(conn, table):
        return
    sql = render_sql(
        """
        INSERT INTO fact_unit_enrolment (
            unit_enrolment_res_key, provider_code, student_res_key, student_identifier,
            chessn, course_admission_res_key, course_of_study_code,
            course_of_study_type, course_res_key, course_code, course_name,
            course_commencement_date, unit_of_study_code, census_date,
            work_experience_code, summer_winter_school_code, discipline_code,
            unit_status_code, mode_of_attendance_code, delivery_location_postcode,
            delivery_location_country_code, student_status_code,
            max_student_contribution_code, year_long_indicator,
            unit_commencement_date, unit_outcome_date, remission_reason_code,
            loan_status, student_loan_res_key, adjusted_loan_amount,
            adjusted_loan_fee, eftsl, amount_charged,
            amount_paid_upfront, help_loan_amount, loan_fee, is_deleted,
            extraction_timestamp, reporting_year
        )
        SELECT
            {{UID16_UnitEnrolmentsResKey}},
            {{E306_ProviderCode}},
            {{UID8_StudentsResKey}},
            {{E313_StudentIdentificationCode}},
            {{E488_CHESSN}},
            {{UID15_CourseAdmissionsResKey}},
            {{E533_CourseOfStudyCode}},
            {{E310_CourseOfStudyType}},
            {{UID5_CoursesResKey}},
            {{E307_CourseCode}},
            {{E308_CourseName}},
            {{E534_CourseOfStudyCommDate}},
            {{E354_UnitOfStudyCode}},
            {{E489_UnitOfStudyCensusDate}},
            {{E337_WorkExperienceInIndustryCode}},
            {{E551_SummerWinterSchoolCode}},
            {{E464_DisciplineCode}},
            {{E355_UnitOfStudyStatusCode}},
            {{E329_ModeOfAttendanceCode}},
            {{E477_DeliveryLocationPostcode}},
            {{E660_DeliveryLocationCountryCode}},
            {{E490_StudentStatusCode}},
            {{E392_MaximumStudentContributionCode}},
            {{E622_UnitOfStudyYearLongIndicator}},
            {{E600_UnitOfStudyCommencementDate}},
            {{E601_UnitOfStudyOutcomeDate}},
            {{E446_RemissionReasonCode}},
            {{A130_LoanStatus}},
            {{UID21_StudentLoansResKey}},
            NULLIF({{E662_AdjustedLoanAmount}}, 'NULL'),
            NULLIF({{E663_AdjustedLoanFee}}, 'NULL'),
            NULLIF({{UE_E339_EFTSL}}, 'NULL'),
            NULLIF({{UE_E384_AmountCharged}}, 'NULL'),
            NULLIF({{UE_E381_AmountPaidUpfront}}, 'NULL'),
            NULLIF({{UE_E558_HELPLoanAmount}}, 'NULL'),
            NULLIF({{UE_E529_LoanFee}}, 'NULL'),
            {{UE_A111_IsDeleted}},
            ?,
            reporting_year
        FROM stg_hep_units_aous
        WHERE extraction_timestamp = ?
          AND {{UID16_UnitEnrolmentsResKey}} IS NOT NULL
        ON CONFLICT(unit_enrolment_res_key) DO UPDATE SET
            provider_code = excluded.provider_code,
            student_res_key = excluded.student_res_key,
            student_identifier = excluded.student_identifier,
            chessn = excluded.chessn,
            course_admission_res_key = excluded.course_admission_res_key,
            course_of_study_code = excluded.course_of_study_code,
            course_of_study_type = excluded.course_of_study_type,
            course_res_key = excluded.course_res_key,
            course_code = excluded.course_code,
            course_name = excluded.course_name,
            course_commencement_date = excluded.course_commencement_date,
            unit_of_study_code = excluded.unit_of_study_code,
            census_date = excluded.census_date,
            work_experience_code = excluded.work_experience_code,
            summer_winter_school_code = excluded.summer_winter_school_code,
            discipline_code = excluded.discipline_code,
            unit_status_code = excluded.unit_status_code,
            mode_of_attendance_code = excluded.mode_of_attendance_code,
            delivery_location_postcode = excluded.delivery_location_postcode,
            delivery_location_country_code = excluded.delivery_location_country_code,
            student_status_code = excluded.student_status_code,
            max_student_contribution_code = excluded.max_student_contribution_code,
            year_long_indicator = excluded.year_long_indicator,
            unit_commencement_date = excluded.unit_commencement_date,
            unit_outcome_date = excluded.unit_outcome_date,
            remission_reason_code = excluded.remission_reason_code,
            loan_status = excluded.loan_status,
            student_loan_res_key = excluded.student_loan_res_key,
            adjusted_loan_amount = excluded.adjusted_loan_amount,
            adjusted_loan_fee = excluded.adjusted_loan_fee,
            eftsl = excluded.eftsl,
            amount_charged = excluded.amount_charged,
            amount_paid_upfront = excluded.amount_paid_upfront,
            help_loan_amount = excluded.help_loan_amount,
            loan_fee = excluded.loan_fee,
            is_deleted = excluded.is_deleted,
            extraction_timestamp = excluded.extraction_timestamp,
            reporting_year = excluded.reporting_year
        """
    )
    conn.execute(sql, (extraction_ts, extraction_ts))


def upsert_unit_enrolment_aou(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_hep_units_aous"
    if not staging_table_exists(conn, table):
        return
    sql = render_sql(
        """
        INSERT INTO fact_unit_enrolment_aou (
            unit_enrolment_aou_res_key, unit_enrolment_res_key, aou_code,
            eftsl, amount_charged, amount_paid_upfront, help_loan_amount,
            loan_fee, is_deleted, extraction_timestamp, reporting_year
        )
        SELECT
            {{UID19_UnitEnrolmentAOUsResKey}},
            {{UID16_UnitEnrolmentsResKey}},
            {{E333_AOUCode}},
            NULLIF({{AOU_E339_EFTSL}}, 'NULL'),
            NULLIF({{AOU_E384_AmountCharged}}, 'NULL'),
            NULLIF({{AOU_E381_AmountPaidUpfront}}, 'NULL'),
            NULLIF({{AOU_E558_HELPLoanAmount}}, 'NULL'),
            NULLIF({{AOU_E529_LoanFee}}, 'NULL'),
            {{AOU_IsDeleted}},
            ?,
            reporting_year
        FROM stg_hep_units_aous
        WHERE extraction_timestamp = ?
          AND {{UID19_UnitEnrolmentAOUsResKey}} IS NOT NULL
        ON CONFLICT(unit_enrolment_aou_res_key) DO UPDATE SET
            unit_enrolment_res_key = excluded.unit_enrolment_res_key,
            aou_code = excluded.aou_code,
            eftsl = excluded.eftsl,
            amount_charged = excluded.amount_charged,
            amount_paid_upfront = excluded.amount_paid_upfront,
            help_loan_amount = excluded.help_loan_amount,
            loan_fee = excluded.loan_fee,
            is_deleted = excluded.is_deleted,
            extraction_timestamp = excluded.extraction_timestamp,
            reporting_year = excluded.reporting_year
        """
    )
    conn.execute(sql, (extraction_ts, extraction_ts))


def upsert_sahelp(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_sahelp"
    if not staging_table_exists(conn, table):
        return
    loan_amount_expr = nullif_or_null(conn, table, "E558_HELPLoanAmount")
    loan_fee_expr = nullif_or_null(conn, table, "E529_LoanFee")
    template = """
        INSERT INTO fact_student_loan (
            student_loan_res_key, provider_res_key, student_res_key,
            course_admission_res_key, loan_type, reporting_year,
            reporting_period, loan_amount, loan_fee, status_code,
            extraction_timestamp
        )
        SELECT
            {{UID21_StudentLoansResKey}},
            {{UID1_ProvidersResKey}},
            {{UID8_StudentsResKey}},
            NULL,
            'SAHELP',
            NULL,
            NULL,
            __LOAN_AMOUNT__,
            __LOAN_FEE__,
            {{A130_LoanStatus}},
            ?
        FROM stg_sahelp
        WHERE extraction_timestamp = ?
          AND {{UID21_StudentLoansResKey}} IS NOT NULL
        ON CONFLICT(student_loan_res_key) DO UPDATE SET
            provider_res_key = excluded.provider_res_key,
            student_res_key = excluded.student_res_key,
            course_admission_res_key = excluded.course_admission_res_key,
            loan_type = excluded.loan_type,
            loan_amount = excluded.loan_amount,
            loan_fee = excluded.loan_fee,
            status_code = excluded.status_code,
            extraction_timestamp = excluded.extraction_timestamp
        """
    template = template.replace("__LOAN_AMOUNT__", loan_amount_expr).replace("__LOAN_FEE__", loan_fee_expr)
    sql = render_sql(template)
    conn.execute(sql, (extraction_ts, extraction_ts))


def upsert_oshelp(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_oshelp"
    if not staging_table_exists(conn, table):
        return
    sql = render_sql(
        """
        INSERT INTO fact_student_loan (
            student_loan_res_key, provider_res_key, student_res_key,
            course_admission_res_key, loan_type, reporting_year,
            reporting_period, loan_amount, loan_fee, status_code,
            extraction_timestamp
        )
        SELECT
            {{UID21_StudentLoansResKey}},
            {{UID1_ProvidersResKey}},
            {{UID8_StudentsResKey}},
            {{UID15_CourseAdmissionsResKey}},
            'OSHELP',
            NULL,
            NULL,
            NULLIF({{E528_OSHELPPaymentAmount}}, 'NULL'),
            NULLIF({{E529_LoanFee}}, 'NULL'),
            {{A130_LoanStatus}},
            ?
        FROM stg_oshelp
        WHERE extraction_timestamp = ?
          AND {{UID21_StudentLoansResKey}} IS NOT NULL
        ON CONFLICT(student_loan_res_key) DO UPDATE SET
            provider_res_key = excluded.provider_res_key,
            student_res_key = excluded.student_res_key,
            course_admission_res_key = excluded.course_admission_res_key,
            loan_type = excluded.loan_type,
            loan_amount = excluded.loan_amount,
            loan_fee = excluded.loan_fee,
            status_code = excluded.status_code,
            extraction_timestamp = excluded.extraction_timestamp
        """
    )
    conn.execute(sql, (extraction_ts, extraction_ts))


def upsert_scholarships(conn: sqlite3.Connection, extraction_ts: str) -> None:
    if staging_table_exists(conn, "stg_commonwealth_scholarships"):
        sql = render_sql(
            """
            INSERT INTO fact_scholarship (
                scholarship_res_key, provider_res_key, student_res_key,
                course_admission_res_key, student_identifier, scholarship_type,
                reporting_year, reporting_period, amount, status_code,
                extraction_timestamp
            )
            SELECT
                {{UID12_StudentCommonwealthScholarshipsResKey}},
                {{UID1_ProvidersResKey}},
                {{UID8_StudentsResKey}},
                NULL,
                {{E313_StudentIdentificationCode}},
                'COMMONWEALTH',
                {{E415_ReportingYear}},
                {{E666_ReportingPeriod}},
                NULLIF({{E598_CommonwealthScholarshipAmount}}, 'NULL'),
                {{E526_CommonwealthScholarshipStatusCode}},
                ?
            FROM stg_commonwealth_scholarships
            WHERE extraction_timestamp = ?
              AND {{UID12_StudentCommonwealthScholarshipsResKey}} IS NOT NULL
            ON CONFLICT(scholarship_res_key) DO UPDATE SET
                provider_res_key = excluded.provider_res_key,
                student_res_key = excluded.student_res_key,
                course_admission_res_key = excluded.course_admission_res_key,
                student_identifier = excluded.student_identifier,
                scholarship_type = excluded.scholarship_type,
                reporting_year = excluded.reporting_year,
                reporting_period = excluded.reporting_period,
                amount = excluded.amount,
                status_code = excluded.status_code,
                extraction_timestamp = excluded.extraction_timestamp
            """
        )
        conn.execute(sql, (extraction_ts, extraction_ts))

    if staging_table_exists(conn, "stg_rtp_scholarships"):
        sql = render_sql(
            """
            INSERT INTO fact_scholarship (
                scholarship_res_key, provider_res_key, student_res_key,
                course_admission_res_key, student_identifier, scholarship_type,
                reporting_year, reporting_period, amount, status_code,
                extraction_timestamp
            )
            SELECT
                {{UID35_RTPScholarshipsResKey}},
                {{UID1_ProvidersResKey}},
                {{UID8_StudentsResKey}},
                {{UID15_CourseAdmissionsResKey}},
                {{E313_StudentIdentificationCode}},
                'RTP_SCHOLARSHIP',
                NULL,
                NULL,
                NULL,
                {{E487_RTPScholarshipType}},
                ?
            FROM stg_rtp_scholarships
            WHERE extraction_timestamp = ?
              AND {{UID35_RTPScholarshipsResKey}} IS NOT NULL
            ON CONFLICT(scholarship_res_key) DO UPDATE SET
                provider_res_key = excluded.provider_res_key,
                student_res_key = excluded.student_res_key,
                course_admission_res_key = excluded.course_admission_res_key,
                student_identifier = excluded.student_identifier,
                scholarship_type = excluded.scholarship_type,
                reporting_year = excluded.reporting_year,
                reporting_period = excluded.reporting_period,
                amount = excluded.amount,
                status_code = excluded.status_code,
                extraction_timestamp = excluded.extraction_timestamp
            """
        )
        conn.execute(sql, (extraction_ts, extraction_ts))

    if staging_table_exists(conn, "stg_rtp_stipend"):
        stipend_student_expr = placeholder_or_null(conn, "stg_rtp_stipend", "UID8_StudentsResKey")
        template = """
            INSERT INTO fact_scholarship (
                scholarship_res_key, provider_res_key, student_res_key,
                course_admission_res_key, student_identifier, scholarship_type,
                reporting_year, reporting_period, amount, status_code,
                extraction_timestamp
            )
            SELECT
                {{UID18_RTPStipendsResKey}},
                {{UID1_ProvidersResKey}},
                __STUDENT__,
                {{UID15_CourseAdmissionsResKey}},
                {{E313_StudentIdentificationCode}},
                'RTP_STIPEND',
                {{E415_ReportingYear}},
                NULL,
                NULLIF({{E623_RTPStipendAmount}}, 'NULL'),
                NULL,
                ?
            FROM stg_rtp_stipend
            WHERE extraction_timestamp = ?
              AND {{UID18_RTPStipendsResKey}} IS NOT NULL
            ON CONFLICT(scholarship_res_key) DO UPDATE SET
                provider_res_key = excluded.provider_res_key,
                student_res_key = excluded.student_res_key,
                course_admission_res_key = excluded.course_admission_res_key,
                student_identifier = excluded.student_identifier,
                scholarship_type = excluded.scholarship_type,
                reporting_year = excluded.reporting_year,
                reporting_period = excluded.reporting_period,
                amount = excluded.amount,
                status_code = excluded.status_code,
                extraction_timestamp = excluded.extraction_timestamp
        """
        sql = render_sql(template.replace("__STUDENT__", stipend_student_expr))
        conn.execute(sql, (extraction_ts, extraction_ts))


def upsert_course_fees(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_campus_course_fees_itsp"
    if not staging_table_exists(conn, table):
        return
    sql = render_sql(
        """
        INSERT INTO fact_course_fee (
            campus_course_fee_res_key, provider_res_key, course_campus_res_key,
            course_res_key, course_code, campus_suburb,
            indicative_student_contribution_csp,
            indicative_tuition_fee_domestic_fp,
            effective_from_date, campus_effective_from_date, extraction_timestamp
        )
        SELECT
            {{UID31_CampusCourseFeesResKey}},
            {{UID1_ProvidersResKey}},
            {{UID4_CoursesOnCampusResKey}},
            {{UID5_CoursesResKey}},
            {{E307_CourseCode}},
            {{E525_CampusSuburb}},
            NULLIF({{E495_IndicativeStudentContributionCSP}}, 'NULL'),
            NULLIF({{E496_IndicativeTuitionFeeDomesticFP}}, 'NULL'),
            {{E609_EffectiveFromDate}},
            {{Campuses_E609_EffectiveFromDate}},
            ?
        FROM stg_campus_course_fees_itsp
        WHERE extraction_timestamp = ?
          AND {{UID31_CampusCourseFeesResKey}} IS NOT NULL
        ON CONFLICT(campus_course_fee_res_key) DO UPDATE SET
            provider_res_key = excluded.provider_res_key,
            course_campus_res_key = excluded.course_campus_res_key,
            course_res_key = excluded.course_res_key,
            course_code = excluded.course_code,
            campus_suburb = excluded.campus_suburb,
            indicative_student_contribution_csp = excluded.indicative_student_contribution_csp,
            indicative_tuition_fee_domestic_fp = excluded.indicative_tuition_fee_domestic_fp,
            effective_from_date = excluded.effective_from_date,
            campus_effective_from_date = excluded.campus_effective_from_date,
            extraction_timestamp = excluded.extraction_timestamp
        """
    )
    conn.execute(sql, (extraction_ts, extraction_ts))


def upsert_aggregated_awards(conn: sqlite3.Connection, extraction_ts: str) -> None:
    table = "stg_aggregated_awards"
    if not staging_table_exists(conn, table):
        return
    sql = render_sql(
        """
        INSERT INTO fact_aggregated_award (
            aggregated_award_res_key, provider_res_key, student_res_key,
            course_res_key, course_code, course_commencement_date,
            course_outcome_code, course_outcome_date,
            hdr_thesis_submission_date, attendance_type_code,
            mode_of_attendance_code, extraction_timestamp
        )
        SELECT
            {{UID47_AggregateAwardsResKey}},
            {{UID1_ProvidersResKey}},
            {{UID8_StudentsResKey}},
            {{UID5_CoursesResKey}},
            {{E307_CourseCode}},
            {{E534_CourseOfStudyCommencementDate}},
            {{E599_CourseOutcomeCode}},
            {{E592_CourseOutcomeDate}},
            {{E591_HDRThesisSubmissionDate}},
            {{E330_AttendanceTypeCode}},
            {{E329_ModeOfAttendanceCode}},
            ?
        FROM stg_aggregated_awards
        WHERE extraction_timestamp = ?
          AND {{UID47_AggregateAwardsResKey}} IS NOT NULL
        ON CONFLICT(aggregated_award_res_key) DO UPDATE SET
            provider_res_key = excluded.provider_res_key,
            student_res_key = excluded.student_res_key,
            course_res_key = excluded.course_res_key,
            course_code = excluded.course_code,
            course_commencement_date = excluded.course_commencement_date,
            course_outcome_code = excluded.course_outcome_code,
            course_outcome_date = excluded.course_outcome_date,
            hdr_thesis_submission_date = excluded.hdr_thesis_submission_date,
            attendance_type_code = excluded.attendance_type_code,
            mode_of_attendance_code = excluded.mode_of_attendance_code,
            extraction_timestamp = excluded.extraction_timestamp
        """
    )
    conn.execute(sql, (extraction_ts, extraction_ts))


UPSERT_SEQUENCE = [
    upsert_dim_provider,
    upsert_dim_student,
    upsert_student_contact,
    upsert_student_citizenship,
    upsert_student_disability,
    upsert_course_of_study,
    upsert_course,
    upsert_course_campus,
    upsert_course_field_of_education,
    upsert_course_admission,
    upsert_course_prior_credit,
    upsert_course_specialisation,
    upsert_unit_enrolment,
    upsert_unit_enrolment_aou,
    upsert_sahelp,
    upsert_oshelp,
    upsert_scholarships,
    upsert_course_fees,
    upsert_aggregated_awards,
]


def parse_extraction_timestamp_from_dir(directory: Path) -> str:
    token = directory.name.split("_")[0]
    for fmt in ("%d%b%Y", "%d%B%Y", "%Y%m%d"):
        try:
            parsed = dt.datetime.strptime(token, fmt)
            break
        except ValueError:
            parsed = None
    if parsed is None:
        parsed = dt.datetime.utcnow()
    return parsed.replace(hour=0, minute=0, second=0, microsecond=0).isoformat(timespec="seconds") + "Z"


def process_directory(conn: sqlite3.Connection, directory: Path) -> Dict[str, int]:
    extraction_ts = parse_extraction_timestamp_from_dir(directory)
    stage_counts: Dict[str, int] = {}
    tabular_files = sorted(
        path
        for path in directory.iterdir()
        if path.is_file() and path.suffix.lower() in SUPPORTED_TABULAR_SUFFIXES
    )
    if not tabular_files:
        print(f"[warn] No CSV/XLSX files found in {directory}")
        return stage_counts
    for file_path in tabular_files:
        suffix = file_path.suffix.lower()
        if suffix == ".csv":
            table, count = load_csv(conn, file_path, extraction_ts)
        elif suffix in {".xlsx", ".xlsm"}:
            table, count = load_excel(conn, file_path, extraction_ts)
        else:  # pragma: no cover - guarded by SUPPORTED_TABULAR_SUFFIXES
            raise ValueError(f"Unsupported file extension: {file_path.suffix}")
        stage_counts[table] = stage_counts.get(table, 0) + count
        print(f"Loaded {count:6d} rows into {table} from {file_path.name}")
    for upserter in UPSERT_SEQUENCE:
        upserter(conn, extraction_ts)
    conn.execute(
        "INSERT INTO etl_load_history (extraction_timestamp, source_directory, notes) VALUES (?, ?, ?)",
        (extraction_ts, str(directory), json.dumps(stage_counts)),
    )
    conn.commit()
    return stage_counts


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "sources",
        nargs="+",
        help="One or more directories containing deidentified extract CSVs",
    )
    parser.add_argument(
        "--db-path",
        default=str(DEFAULT_DB_PATH),
        help=f"SQLite database path (default: {DEFAULT_DB_PATH})",
    )
    args = parser.parse_args()

    conn = sqlite3.connect(Path(args.db_path))
    conn.execute("PRAGMA foreign_keys = ON")
    apply_schema(conn)

    for source in args.sources:
        directory = Path(source)
        if not directory.is_dir():
            print(f"[warn] Skipping {directory} (not a directory)")
            continue
        print(f"\nProcessing directory: {directory}")
        counts = process_directory(conn, directory)
        print(f"Summary for {directory.name}: {json.dumps(counts, indent=2)}")

    conn.close()


if __name__ == "__main__":
    main()
